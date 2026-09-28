"""Malaka oshirish — uchidan-uchiga tekshiruv HAQIQIY bazada, lekin hech narsa saqlanmaydi.

Hamma yozuv bitta tashqi tranzaksiya ichida qilinadi (marshrutlarning
commit'lari savepoint bo'lib qoladi) va oxirida butunlay bekor qilinadi.
Oxirgi bo'lim OpenAI'ga ikkita kichik so'rov yuboradi (fayldan test importi).

Ishlatish (serverda, /home/imentor):
    docker compose -f docker-compose.prod.yml run --rm --no-deps -T \
        -v "$PWD/deploy/malaka/smoke_malaka.py:/smoke.py" --entrypoint "" backend_fastapi \
        sh -c "pip install -q httpx==0.27.2; python /smoke.py"
"""

import datetime as dt
import io
import sys

sys.path.insert(0, "/app")

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy.orm import Session  # noqa: E402

from app.core.db import engine, get_db  # noqa: E402
from app.core.security import decode_token  # noqa: E402
from app.main import app  # noqa: E402
from app.models.online_edu import (  # noqa: E402
    MalakaListener,
    OnlineGroup,
    OnlineGroupCourse,
    OnlineSyllabus,
    OnlineTeacher,
    OnlineTeacherCourse,
)
from app.services import auth_service  # noqa: E402
from app.services import password_policy_service as pwd  # noqa: E402

conn = engine.connect()
outer = conn.begin()


def session() -> Session:
    return Session(bind=conn, join_transaction_mode="create_savepoint")


def override():
    s = session()
    try:
        yield s
    finally:
        s.close()


app.dependency_overrides[get_db] = override
c = TestClient(app)
passed = 0


def check(cond, msg):
    global passed
    if not cond:
        raise SystemExit(f"  XATO: {msg}")
    passed += 1
    print("  ok ", msg)


def build(model, **fields):
    if hasattr(model, "created_at") and "created_at" not in fields:
        fields["created_at"] = dt.datetime.now(dt.timezone.utc)
    return model(**fields)


def Q(n, prefix=True):
    opts = ["A) to'g'ri", "B) xato", "C) xato"] if prefix else ["to'g'ri", "xato", "xato"]
    return [{"question": f"Savol {i + 1}?", "options": opts, "correctOptionIndex": 0, "explanation": ""} for i in range(n)]


now = dt.datetime.now(dt.timezone.utc)
try:
    db = Session(bind=conn, join_transaction_mode="create_savepoint", expire_on_commit=False)
    common = dict(department_name="", department_id=None, description="", instruction_language="uz",
                  file_name="", variants=[], name_i18n={}, topics_i18n={}, sort_order=99, is_active=True,
                  created_at=now, updated_at=now)
    syl = OnlineSyllabus(subject_name="SMOKE malaka", subject_code="smoke-malaka-zz", topics=[], program="malaka", **common)
    osyl = OnlineSyllabus(subject_name="SMOKE online", subject_code="smoke-online-zz",
                          topics=[{"code": "1", "title": "X", "type": "lecture"}], program="online", **common)
    db.add_all([syl, osyl])
    db.flush()
    grp = OnlineGroup(name="SMOKE malaka guruh", program="malaka", is_active=True, created_at=now)
    db.add(grp)
    db.flush()
    db.add(build(OnlineGroupCourse, group_id=grp.id, syllabus_id=syl.id, variant_label=""))
    auth_service.create_user(db, "SMOKETEACHER1", "boshlangich1", "Test", "Oqituvchi")
    pwd.require_change(db, "SMOKETEACHER1")
    teacher = OnlineTeacher(owner_key="SMOKETEACHER1", full_name="Test", is_active=True, created_at=now)
    db.add(teacher)
    db.flush()
    db.add(build(OnlineTeacherCourse, teacher_id=teacher.id, syllabus_id=syl.id, variant_label=""))
    db.add(build(OnlineTeacherCourse, teacher_id=teacher.id, syllabus_id=osyl.id, variant_label=""))
    lu = auth_service.create_user(db, "ZZ9999999", "ZZ9999999", "Sinov", "Tinglovchi")
    auth_service.set_user_role_group(db, lu, "student")
    pwd.require_change(db, "ZZ9999999")
    db.add(MalakaListener(username="ZZ9999999", full_name="Tinglovchi Sinov", passport="ZZ 9999999",
                          group_id=grp.id, is_active=True, created_at=now))
    db.commit()
    db.close()
    S, O = syl.id, osyl.id
    base = {"syllabus_id": S, "variant_label": ""}
    topics_url = f"/api/v1/online/student/topics/?syllabus_id={S}&variant_label="

    print("1) Tinglovchi kirishi va majburiy parol")
    check(c.post("/api/v1/auth/malaka-login/", json={"login": "ZZ9999999", "password": "nimadir"}).status_code == 401,
          "noto'g'ri parol -> 401")
    r = c.post("/api/v1/auth/malaka-login/", json={"login": "zz 9999999", "password": "zz 9999999"})
    check(r.status_code == 200, "pasport (kichik harf, bo'sh joy bilan) -> kirdi")
    L = r.json()
    check(L["must_change_password"] and L["group_name"] == "SMOKE malaka guruh" and L["student_id"] == "ZZ9999999",
          "javob: parol almashtirish talab, guruh va ID to'g'ri")
    check(decode_token(L["access"]).get("mcp") == 1, "tokenda 'mcp' belgisi bor")
    hL = {"Authorization": f"Bearer {L['access']}"}
    r = c.get("/api/v1/online/student/subjects/", headers=hL)
    check(r.status_code == 403 and "parol" in r.json()["detail"].lower(), "parol almashtirilmaguncha portal yopiq")
    check(c.post("/api/v1/auth/change-password/", headers=hL,
                 json={"current_password": "ZZ 9999999", "new_password": "zz-9999999"}).status_code == 400,
          "login bilan bir xil parol rad etiladi")
    r = c.post("/api/v1/auth/change-password/", headers=hL,
               json={"current_password": "ZZ 9999999", "new_password": "Yangi-parol-77"})
    check(r.status_code == 200, f"parol almashtirildi ({r.status_code})")
    check(c.get("/api/v1/online/student/subjects/", headers=hL).status_code == 200, "o'sha token darhol ishlaydi")
    r = c.post("/api/v1/auth/token/refresh/", json={"refresh": L["refresh"]})
    check(r.status_code == 200 and "mcp" not in decode_token(r.json()["access"]), "yangilangan tokenda belgi yo'q")
    check(c.post("/api/v1/auth/malaka-login/", json={"login": "ZZ9999999", "password": "ZZ9999999"}).status_code == 401,
          "eski (pasport) parol endi ishlamaydi")
    check(c.post("/api/v1/auth/malaka-login/", json={"login": "zz9999999", "password": "Yangi-parol-77"}).status_code == 200,
          "yangi parol bilan kiradi")

    print("2) O'qituvchi")
    r = c.post("/api/v1/auth/local-login/", json={"phone_digits": "smoketeacher1", "password": "boshlangich1"})
    check(r.status_code == 200 and r.json()["must_change_password"], "o'qituvchi (kichik harfli login) kirdi, parol talab")
    hT = {"Authorization": f"Bearer {r.json()['access']}"}
    check(c.get("/api/v1/online/teacher/me/?program=malaka", headers=hT).status_code == 403, "parolsiz kabinet yopiq")
    check(c.post("/api/v1/auth/change-password/", headers=hT,
                 json={"current_password": "boshlangich1", "new_password": "Oqituvchi-99"}).status_code == 200,
          "o'qituvchi parolini almashtirdi")
    me = c.get("/api/v1/online/teacher/me/?program=malaka", headers=hT).json()
    check([x["syllabus_id"] for x in me["courses"]] == [S] and me["courses"][0]["program"] == "malaka",
          "malaka portalida faqat malaka fani")
    me = c.get("/api/v1/online/teacher/me/?program=online", headers=hT).json()
    check([x["syllabus_id"] for x in me["courses"]] == [O], "online portalida faqat online fan")

    print("3) Mavzular")
    for title in ["Pedagogika asoslari", "Tibbiy ta'limda baholash", "O'chiriladigan"]:
        r = c.post("/api/v1/malaka/teacher/topics/", headers=hT, json={"syllabus_id": S, "title": title})
        check(r.status_code == 201, f"mavzu qo'shildi: {title} (kod {r.json().get('code')})")
    check(c.post("/api/v1/malaka/teacher/topics/", headers=hT,
                 json={"syllabus_id": S, "title": "pedagogika  ASOSLARI"}).status_code == 409, "takroriy nom rad etiladi")
    r = c.patch("/api/v1/malaka/teacher/topics/2/", headers=hT, json={"syllabus_id": S, "title": "Tibbiy ta'limda nazorat"})
    check(r.status_code == 200 and r.json()["title"] == "Tibbiy ta'limda nazorat", "nomi o'zgartirildi")
    r = c.post("/api/v1/malaka/teacher/topics/2/move/", headers=hT, json={"syllabus_id": S, "direction": "up"})
    check([t["code"] for t in r.json()] == ["2", "1", "3"], "tartib o'zgardi")
    check(c.delete(f"/api/v1/malaka/teacher/topics/3/?syllabus_id={S}", headers=hT).status_code == 204, "bo'sh mavzu o'chirildi")
    check(c.post("/api/v1/malaka/teacher/topics/", headers=hT, json={"syllabus_id": O, "title": "Online mavzu"}).status_code == 400,
          "online fanga mavzu qo'shib bo'lmaydi")
    tt = c.get(f"/api/v1/online/teacher/topics/?syllabus_id={S}&variant_label=", headers=hT).json()
    check(set(tt[0]["has"]) == {"lecture", "presentation", "video", "practical", "test"},
          "malaka turlari: tarqatma va masala yo'q, amaliy bor")

    print("4) Materiallar")
    mat = "/api/v1/online/teacher/materials/"
    check(c.post(mat, headers=hT, json={**base, "topic_code": "1", "kind": "handout", "external_url": "https://x.uz"}).status_code == 400,
          "tarqatma malakada rad etiladi")
    check(c.post(mat, headers=hT, json={**base, "syllabus_id": O, "topic_code": "1", "kind": "practical",
                                        "payload": {"questions": Q(2)}}).status_code == 400, "amaliy online fanda rad etiladi")
    check(c.post(mat, headers=hT, json={**base, "topic_code": "99", "kind": "lecture", "payload": {"text": "x"}}).status_code == 400,
          "yo'q mavzuga material rad etiladi")
    check(c.post(mat, headers=hT, json={**base, "topic_code": "1", "kind": "test",
                                        "payload": {"questions": [{"question": "Q", "options": ["bitta"], "correctOptionIndex": 0}]}}).status_code == 400,
          "chala savol rad etiladi")
    check(c.post(mat, headers=hT, json={**base, "topic_code": "1", "kind": "lecture", "payload": {"text": "Ma'ruza"}}).status_code == 201,
          "ma'ruza saqlandi")
    r = c.post(mat, headers=hT, json={**base, "topic_code": "1", "kind": "test", "payload": {"questions": Q(5)}})
    check(r.status_code == 201 and r.json()["payload"]["questions"][0]["options"] == ["to'g'ri", "xato", "xato"],
          "test saqlandi, 'A)' belgilari tozalandi")
    check(c.post(mat, headers=hT, json={**base, "topic_code": "1", "kind": "practical", "payload": {"questions": Q(3)}}).status_code == 201,
          "amaliy mashg'ulot saqlandi")
    r = c.post(mat, headers=hT, json={**base, "topic_code": "__entry__", "kind": "test", "payload": {"questions": Q(4)}})
    check(r.status_code == 201 and r.json()["payload"]["published"] is True, "kirish testi darhol e'lon qilindi")
    r = c.post(mat, headers=hT, json={**base, "topic_code": "__exit__", "kind": "test", "payload": {"questions": Q(4)}})
    check(r.status_code == 201 and r.json()["payload"]["published"] is False, "chiqish testi e'lon qilinmay saqlandi")

    print("5) Kirish testi qulfi")
    s = [x for x in c.get("/api/v1/online/student/subjects/", headers=hL).json() if x["syllabus_id"] == S][0]
    check(s["entry"]["available"] and s["entry"]["used"] == 0 and not s["exit"]["available"],
          "kirish testi ochiq, chiqish testi yashirin")
    tp = c.get(topics_url, headers=hL).json()
    check(all(not t["is_open"] and t["locked_reason"] == "entry" for t in tp), "kirish testisiz hamma mavzu yopiq")
    check(c.get(f"/api/v1/online/student/topic/?syllabus_id={S}&variant_label=&topic_code=1", headers=hL).status_code == 403,
          "mavzu materiali berilmaydi")
    et = c.get(f"/api/v1/malaka/student/subject-test/?syllabus_id={S}&code=entry", headers=hL).json()
    check(len(et["questions"]) == 4 and "correct_index" not in et["questions"][0], "kirish testi javobsiz keladi")
    check(c.get(f"/api/v1/malaka/student/subject-test/?syllabus_id={S}&code=exit", headers=hL).status_code == 404,
          "e'lon qilinmagan chiqish testi berilmaydi")
    att = "/api/v1/malaka/student/attempt/"
    r = c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "__entry__", "kind": "entry", "answers": [0, 1, 0, 1]})
    check(r.status_code == 200 and r.json()["score"] == 2 and r.json()["attempts"]["finished"], "kirish testi: 2/4, urinish tugadi")
    check("correct_index" in r.json()["questions"][0], "yagona urinishdan keyin javoblar ochildi")
    check(c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "__entry__", "kind": "entry",
                                        "answers": [0, 0, 0, 0]}).status_code == 409, "kirish testi qayta topshirilmaydi")

    print("6) Mavzu va 4 ta urinish")
    tp = {t["topic_code"]: t for t in c.get(topics_url, headers=hL).json()}
    check(tp["1"]["is_open"] and not tp["2"]["is_open"] and tp["2"]["locked_reason"] == "empty",
          "materiali bor mavzu ochildi, bo'shi yopiq")
    d = c.get(f"/api/v1/online/student/topic/?syllabus_id={S}&variant_label=&topic_code=1", headers=hL).json()
    check({m["kind"] for m in d["materials"]} == {"lecture", "test", "practical"}, "mavzu bo'limlari keldi")
    tm = [m for m in d["materials"] if m["kind"] == "test"][0]
    check("correct_index" not in tm["questions"][0] and tm["attempts"]["max_attempts"] == 4, "test javobsiz, 4 urinish")
    for i, ans in enumerate([[1] * 5, [0, 1, 1, 1, 1], [0, 0, 1, 1, 1], [1] * 5], 1):
        r = c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "1", "kind": "test", "answers": ans})
        check(r.status_code == 200, f"{i}-urinish: {r.json()['score']}/5")
        if i < 4:
            check("correct_index" not in r.json()["questions"][0], f"{i}-urinishdan keyin javoblar hali yopiq")
    check("correct_index" in r.json()["questions"][0] and r.json()["attempts"]["best_percent"] == 40,
          "4-urinishdan keyin javoblar ochildi, eng yaxshisi 40%")
    check(c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "1", "kind": "test", "answers": [0] * 5}).status_code == 409,
          "5-urinish rad etiladi")
    check(c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "1", "kind": "practical", "answers": [0]}).status_code == 409,
          "javoblar soni mos emas -> rad (test o'zgargan)")
    r = c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "1", "kind": "practical", "answers": [0, 0, 0]})
    check(r.status_code == 200 and "correct_index" in r.json()["questions"][0], "amaliy 100% -> javoblar darhol ochildi")
    check(c.post("/api/v1/online/student/view/", headers=hL, json={**base, "topic_code": "1", "kind": "lecture"}).status_code == 200,
          "ma'ruza 'ko'rildi' deb belgilandi")
    check(c.post("/api/v1/online/student/test/", headers=hL, json={**base, "topic_code": "1", "answers": [0]}).status_code == 400,
          "online test yo'li malakada yopiq")
    tp = {t["topic_code"]: t for t in c.get(topics_url, headers=hL).json()}
    check(tp["1"]["grade"] == 70, "mavzu bahosi = (40 + 100) / 2 = 70")

    print("7) Chiqish testi")
    check(c.post(mat, headers=hT, json={**base, "topic_code": "__exit__", "kind": "test",
                                        "payload": {"questions": Q(4), "published": True}}).status_code == 201,
          "o'qituvchi chiqish testini e'lon qildi")
    r = c.post(att, headers=hL, json={"syllabus_id": S, "topic_code": "__exit__", "kind": "exit", "answers": [0, 0, 0, 0]})
    check(r.status_code == 200 and r.json()["score"] == 4, "chiqish testi: 4/4")

    print("8) O'qituvchi natijalari")
    p = c.get(f"/api/v1/malaka/teacher/progress/?syllabus_id={S}", headers=hT).json()
    st = [x for x in p["students"] if x["student_id"] == "ZZ9999999"][0]
    check(st["entry"]["best_percent"] == 50 and st["exit"]["best_percent"] == 100 and st["average_grade"] == 70,
          "kirish 50%, chiqish 100%, o'rtacha baho 70")
    t1 = [t for t in st["topics"] if t["topic_code"] == "1"][0]
    check(t1["viewed"]["lecture"] and t1["test"]["used"] == 4, "ko'rilgan ma'ruza va urinishlar ko'rinadi")
    check(p["groups"] == [{"id": grp.id, "name": "SMOKE malaka guruh", "count": 1}], "guruh va tinglovchilar soni")
    check(c.delete(f"/api/v1/malaka/teacher/topics/1/?syllabus_id={S}", headers=hT).status_code == 409,
          "materiali bor mavzu o'chirilmaydi")
    check(c.get(f"/api/v1/malaka/teacher/progress/?syllabus_id={S}", headers=hL).status_code == 403,
          "tinglovchi o'qituvchi hisobotini ko'ra olmaydi")

    print("9) Fayldan test — AI")
    text = (
        "1. Kattalarda arterial bosimning me'yori qaysi?\nA) 120/80 mm sim. ust.\nB) 200/150 mm sim. ust.\n"
        "C) 60/30 mm sim. ust.\nJavob: A\n\n2. Sog'lom odamda tana haroratining me'yori?\nA) 36,6 °C\nB) 40 °C\nC) 33 °C\n"
    )
    r = c.post("/api/v1/malaka/teacher/test-import/", headers=hT, data={"syllabus_id": str(S), "text": text})
    check(r.status_code == 200 and len(r.json()["questions"]) == 2, f"matndan 2 savol ajratildi ({r.status_code})")
    qs = r.json()["questions"]
    check(qs[0]["answer_source"] == "document" and qs[0]["correctOptionIndex"] == 0, "1-savol: javob fayldan olindi")
    check(qs[1]["answer_source"] == "ai" and qs[1]["correctOptionIndex"] == 0, "2-savol: faylda javob yo'q — AI to'g'ri tanladi")
    from docx import Document

    doc = Document()
    doc.add_paragraph("1. Yurak nechta kameradan iborat?")
    doc.add_paragraph("A) 2 ta")
    doc.add_paragraph().add_run("B) 4 ta").bold = True
    doc.add_paragraph("C) 3 ta")
    buf = io.BytesIO()
    doc.save(buf)
    r = c.post("/api/v1/malaka/teacher/test-import/", headers=hT, data={"syllabus_id": str(S)},
               files={"file": ("test.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    qs = r.json().get("questions", [])
    check(r.status_code == 200 and len(qs) == 1 and qs[0]["correctOptionIndex"] == 1 and qs[0]["options"][1] == "4 ta",
          "Word: qalin yozilgan javob to'g'ri deb olindi, 'B)' tozalandi")
    check(c.post("/api/v1/malaka/teacher/test-import/", headers=hT, data={"syllabus_id": str(S)},
                 files={"file": ("x.xlsx", b"abc", "application/octet-stream")}).status_code == 400,
          "noto'g'ri fayl turi rad etiladi")

    print(f"\nHAMMASI O'TDI: {passed} ta tekshiruv")
finally:
    outer.rollback()
    conn.close()
    print("Bazaga hech narsa yozilmadi — tranzaksiya bekor qilindi.")
