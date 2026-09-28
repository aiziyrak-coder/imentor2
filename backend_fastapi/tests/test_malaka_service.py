"""Malaka oshirish qoidalari — urinishlar, baho, mavzu kodi va savollar.

Bular tinglovchi NIMA ko'rishini, necha marta topshira olishini va qanday
baho olishini hal qiladi, shuning uchun bazasiz, alohida sinaladi.
"""

import datetime as dt
import io
from types import SimpleNamespace

import pytest

from app.core.staff_login import normalize_listener_login
from app.models.online_edu import OnlineSyllabus
from app.services import malaka_service as mk


def syllabus(program="malaka", topics=None):
    s = OnlineSyllabus()
    s.program = program
    s.topics = topics if topics is not None else []
    s.variants = []
    return s


def attempt(no, score, total=10):
    return SimpleNamespace(
        attempt_no=no,
        score=score,
        total=total,
        answers=[0] * total,
        submitted_at=dt.datetime(2026, 9, 14, 9, no, tzinfo=dt.timezone.utc),
    )


def q(question="Savol?", options=("Bir", "Ikki", "Uch"), correct=0, **extra):
    return {"question": question, "options": list(options), "correctOptionIndex": correct, **extra}


class TestLogin:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("AA 1234567", "AA1234567"),
            ("aa1234567", "AA1234567"),
            (" aa-1234567 ", "AA1234567"),
            ("2062208", "2062208"),
            # Kirill klaviaturasida terilgan "АВ" — ko'zga bir xil.
            ("АВ 1000002", "AB1000002"),
            ("", ""),
        ],
    )
    def test_normalize(self, raw, expected):
        assert normalize_listener_login(raw) == expected


class TestProgram:
    def test_malaka_turlari(self):
        assert mk.kinds_for(syllabus("malaka")) == (
            "lecture", "presentation", "video", "practical", "test",
        )

    def test_online_turlari_ozgarmagan(self):
        kinds = mk.kinds_for(syllabus("online"))
        assert "practical" not in kinds
        assert {"case", "handout"} <= set(kinds)

    def test_dasturi_yozilmagan_eski_qator(self):
        s = syllabus()
        s.program = None
        assert not mk.is_malaka(s)
        assert mk.kinds_for(s) == mk.kinds_for(syllabus("online"))

    def test_urinish_turi(self):
        assert mk.attempt_kind(mk.ENTRY_CODE, "test") == "entry"
        assert mk.attempt_kind(mk.EXIT_CODE, "test") == "exit"
        assert mk.attempt_kind("3", "practical") == "practical"
        assert mk.MAX_ATTEMPTS == {"test": 4, "practical": 4, "entry": 1, "exit": 1}


class TestSummarize:
    def test_urinishsiz(self):
        s = mk.summarize([], "test")
        assert (s["used"], s["left"], s["finished"], s["best_percent"]) == (0, 4, False, None)
        assert not mk.reveal_answers(s)

    def test_eng_yaxshisi_hisoblanadi(self):
        s = mk.summarize([attempt(1, 4), attempt(2, 9), attempt(3, 6)], "test")
        assert s["best_percent"] == 90
        assert s["last_score"] == 6
        assert s["left"] == 1
        # Urinish qolgan — javoblar hali yopiq.
        assert not mk.reveal_answers(s)

    def test_urinishlar_tugagach_javob_ochiladi(self):
        s = mk.summarize([attempt(i, 5) for i in range(1, 5)], "practical")
        assert s["finished"] and s["left"] == 0
        assert mk.reveal_answers(s)

    def test_kirish_testi_bir_urinish(self):
        s = mk.summarize([attempt(1, 3)], "entry")
        assert s["finished"] and s["max_attempts"] == 1

    def test_100_foizda_javob_ochiladi(self):
        s = mk.summarize([attempt(1, 10)], "test")
        assert not s["finished"]
        assert mk.reveal_answers(s)

    def test_nol_savolli_urinish_yiqitmaydi(self):
        assert mk.summarize([attempt(1, 0, total=0)], "test")["best_percent"] == 0


class TestTopicResult:
    def test_bittasi_topshirilgan_baho_yoq(self):
        r = mk.topic_result({("1", "test"): [attempt(1, 8)]}, "1", {"test", "practical"})
        assert r["grade"] is None
        assert r["test"]["best_percent"] == 80
        assert r["practical"]["used"] == 0

    def test_ikkalasi_topshirilgan(self):
        attempts = {("1", "test"): [attempt(1, 8)], ("1", "practical"): [attempt(1, 5), attempt(2, 6)]}
        assert mk.topic_result(attempts, "1", {"test", "practical", "lecture"})["grade"] == 70

    def test_faqat_test_bor(self):
        r = mk.topic_result({("1", "test"): [attempt(1, 8)]}, "1", {"test", "lecture"})
        assert r["grade"] == 80
        assert r["practical"] is None

    def test_testsiz_mavzu(self):
        r = mk.topic_result({}, "1", {"lecture", "video"})
        assert r == {"test": None, "practical": None, "grade": None}

    def test_boshqa_mavzu_aralashmaydi(self):
        r = mk.topic_result({("2", "test"): [attempt(1, 10)]}, "1", {"test"})
        assert r["grade"] is None


class TestTopicCodes:
    def test_keyingi_kod(self):
        assert mk.next_topic_code(syllabus(topics=[{"code": "1"}, {"code": "2"}])) == "3"

    def test_bosh_fan(self):
        assert mk.next_topic_code(syllabus()) == "1"

    def test_ochirilgan_kod_qayta_berilmaydi(self):
        s = syllabus(topics=[{"code": "1"}])
        assert mk.next_topic_code(s, {"1", "2", "3", mk.ENTRY_CODE}) == "4"


class TestQuestions:
    def test_toza_savol(self):
        out = mk.clean_questions([q(explanation="  izoh ")])
        assert out[0]["options"] == ["Bir", "Ikki", "Uch"]
        assert out[0]["explanation"] == "izoh"
        assert "answer_source" not in out[0]

    def test_harf_belgilari_olib_tashlanadi(self):
        out = mk.clean_questions([q(options=["A) Bir", "B) Ikki", "C) Uch"])])
        assert out[0]["options"] == ["Bir", "Ikki", "Uch"]

    def test_bittasida_harf_bolsa_tegilmaydi(self):
        options = ["E. coli", "Stafilokokk", "Streptokokk"]
        assert mk.clean_questions([q(options=options)])[0]["options"] == options

    def test_bosh_variant_indeksni_buzmaydi(self):
        out = mk.clean_questions([q(options=["Bir", "", "Uch"], correct=2)])
        assert out[0]["options"] == ["Bir", "Uch"]
        assert out[0]["correctOptionIndex"] == 1

    def test_togri_javob_bosh_variantda(self):
        with pytest.raises(ValueError):
            mk.clean_questions([q(options=["Bir", "", "Uch"], correct=1)])

    def test_savol_raqami(self):
        assert mk.clean_questions([q(question="12. Nima?")])[0]["question"] == "Nima?"
        assert mk.clean_questions([q(question="2.5 mg dozada nima?")])[0]["question"].startswith("2.5")

    def test_qaysi_savol_xato_ekani_aytiladi(self):
        with pytest.raises(ValueError, match="2-savol"):
            mk.clean_questions([q(), q(options=["Yolg'iz"])])

    def test_bosh_royxat(self):
        with pytest.raises(ValueError):
            mk.clean_questions([])

    def test_chegara(self):
        with pytest.raises(ValueError):
            mk.clean_questions([q()] * (mk.MAX_QUESTIONS + 1))


class TestValidateMaterial:
    def topics(self):
        return syllabus(topics=[{"code": "1", "title": "Pedagogika asoslari"}])

    def test_online_da_amaliy_yoq(self):
        with pytest.raises(ValueError):
            mk.validate_material(syllabus("online"), "1", "practical", {})

    def test_online_payload_tegilmaydi(self):
        assert mk.validate_material(syllabus("online"), "1", "lecture", {"text": "x"}) is None

    def test_malakada_tarqatma_va_masala_yoq(self):
        for kind in ("handout", "case"):
            with pytest.raises(ValueError):
                mk.validate_material(self.topics(), "1", kind, {})

    def test_mavzu_topilmadi(self):
        with pytest.raises(ValueError):
            mk.validate_material(self.topics(), "9", "lecture", {"text": "x"})

    def test_kirish_testiga_faqat_test(self):
        with pytest.raises(ValueError):
            mk.validate_material(self.topics(), mk.ENTRY_CODE, "lecture", {"text": "x"})

    def test_kirish_ochiq_chiqish_yopiq(self):
        s = self.topics()
        assert mk.validate_material(s, mk.ENTRY_CODE, "test", {"questions": [q()]})["published"] is True
        assert mk.validate_material(s, mk.EXIT_CODE, "test", {"questions": [q()]})["published"] is False
        assert mk.validate_material(
            s, mk.EXIT_CODE, "test", {"questions": [q()], "published": True}
        )["published"] is True

    def test_amaliy_tozalanadi(self):
        out = mk.validate_material(
            self.topics(), "1", "practical",
            {"questions": [q(options=["A) x", "B) y"])], "ortiqcha": 1},
        )
        assert out == {"questions": [{
            "question": "Savol?", "options": ["x", "y"], "correctOptionIndex": 0, "explanation": "",
        }]}


class TestChunks:
    def test_savol_ortasidan_kesilmaydi(self):
        text = "\n".join(f"{i}. Savol {i}?\nA) bir\nB) ikki\nC) uch" for i in range(1, 300))
        chunks = mk._chunks(text, size=500)
        assert len(chunks) > 1
        for chunk in chunks:
            assert mk._QUESTION_START.match(chunk.split("\n")[0])
        assert "\n".join(chunks) == text


class TestExtract:
    def test_qalin_qator_belgilanadi(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("1. Savol?")
        doc.add_paragraph().add_run("A) To'g'ri").bold = True
        doc.add_paragraph("B) Noto'g'ri")
        buf = io.BytesIO()
        doc.save(buf)
        lines = mk.extract_text("test.docx", buf.getvalue()).split("\n")
        assert not lines[0].startswith(mk.EMPHASIS_MARK)
        assert lines[1] == mk.EMPHASIS_MARK + "A) To'g'ri"
        assert not lines[2].startswith(mk.EMPHASIS_MARK)

    def test_matn_fayli(self):
        assert mk.extract_text("t.txt", "1. Savol?".encode("utf-8")) == "1. Savol?"

    def test_notanish_tur(self):
        with pytest.raises(ValueError):
            mk.extract_text("t.xlsx", b"x")


class TestLockKey:
    def test_barqaror_va_farqli(self):
        a = mk.attempt_lock_key("AA1234567", 5, "1", "test")
        assert a == mk.attempt_lock_key("AA1234567", 5, "1", "test")
        assert a != mk.attempt_lock_key("AA1234567", 5, "1", "practical")
        assert -(2**63) <= a < 2**63
